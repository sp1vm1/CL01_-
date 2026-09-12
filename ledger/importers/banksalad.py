"""뱅크샐러드 '내보내기' 엑셀 가져오기.

- '가계부 내역' 시트: 날짜/시간/타입/대분류/소분류/내용/금액/화폐/결제수단/메모
- '뱅샐현황' 시트 3.재무현황: 내보낸 시점의 계좌별 잔액
"""
from __future__ import annotations

import hashlib
from datetime import datetime, time as dtime
from pathlib import Path

from openpyxl import load_workbook

from ..models import Account, Transaction

KIND_BY_SECTION = {
    "자유입출금 자산": "deposit", "저축성 자산": "savings", "전자금융 자산": "pay",
    "투자성 자산": "investment", "연금 자산": "pension", "기타 실물 자산": "other",
    "현금 자산": "other", "보험 자산": "other", "신탁 자산": "investment",
}


def account_id_for(payment_method: str) -> str:
    return "banksalad:" + hashlib.sha1(payment_method.encode("utf-8")).hexdigest()[:12]


def _fmt_date(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10]


def _fmt_time(v) -> str:
    if isinstance(v, (datetime, dtime)):
        return v.strftime("%H:%M:%S")
    return str(v or "")


def read_transactions(path: str | Path) -> list[Transaction]:
    wb = load_workbook(path, data_only=True)
    if "가계부 내역" not in wb.sheetnames:
        raise ValueError("'가계부 내역' 시트가 없습니다. 뱅크샐러드 내보내기 파일인지 확인하세요.")
    ws = wb["가계부 내역"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    idx = {h: i for i, h in enumerate(header)}
    need = ["날짜", "타입", "대분류", "내용", "금액", "결제수단"]
    missing = [c for c in need if c not in idx]
    if missing:
        raise ValueError(f"필수 컬럼 누락: {missing}")
    out: list[Transaction] = []
    for r in rows:
        if r is None or r[idx["날짜"]] is None:
            continue
        method = str(r[idx["결제수단"]] or "미지정")
        category = str(r[idx["대분류"]] or "")
        tx_type = str(r[idx["타입"]] or "")
        if tx_type == "이체" and category in ("내계좌이체", "카드대금", "현금", "이체"):
            category = "이체"
        out.append(Transaction(
            account_id=account_id_for(method),
            date=_fmt_date(r[idx["날짜"]]),
            time=_fmt_time(r[idx.get("시간", -1)] if "시간" in idx else ""),
            description=str(r[idx["내용"]] or ""),
            amount=int(r[idx["금액"]] or 0),
            category=category,
            subcategory=str(r[idx["소분류"]] or "") if "소분류" in idx else "",
            memo=str(r[idx["메모"]] or "") if "메모" in idx and r[idx["메모"]] else "",
            source="banksalad",
        ))
    return out


def read_accounts(path: str | Path) -> list[Account]:
    """재무현황 표를 읽어 계좌 스냅샷을 만든다. 결제수단명과 상품명이 같으면 같은 계좌 id 를 쓴다."""
    wb = load_workbook(path, data_only=True)
    if "뱅샐현황" not in wb.sheetnames:
        return []
    ws = wb["뱅샐현황"]
    accounts: list[Account] = []
    section = ""
    in_assets = False
    col_item = col_name = col_amount = None
    seen: dict[str, int] = {}
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        if any(v == "3.재무현황" for v in vals):
            in_assets = True
            continue
        if in_assets and any(v == "4.보험현황" for v in vals):
            break
        if not in_assets:
            continue
        if col_item is None:
            # 헤더 행("항목", "상품명", "금액")에서 열 위치를 찾는다. 자산 표가 왼쪽에 먼저 온다.
            if "항목" in vals and "상품명" in vals and "금액" in vals:
                col_item, col_name, col_amount = vals.index("항목"), vals.index("상품명"), vals.index("금액")
            continue
        item = vals[col_item] if col_item < len(vals) else None
        name = vals[col_name] if col_name < len(vals) else None
        amount = vals[col_amount] if col_amount < len(vals) else None
        if isinstance(item, str) and item in KIND_BY_SECTION:
            section = item
        if not section or not isinstance(name, str) or not isinstance(amount, (int, float)):
            continue
        if item in ("항목", "총자산"):
            continue
        seen[name] = seen.get(name, 0) + 1
        display = name if seen[name] == 1 else f"{name} ({seen[name]})"
        accounts.append(Account(
            id=account_id_for(name) if seen[name] == 1 else account_id_for(display),
            provider="banksalad", bank="뱅크샐러드", name=display, number="",
            kind=KIND_BY_SECTION[section], balance=int(round(amount)),
        ))
    return accounts

